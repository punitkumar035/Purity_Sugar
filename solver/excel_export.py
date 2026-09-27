"""
solver/excel_export.py
Purity for Sugar — Excel Report Generation Module
Produces formatted Excel workbooks (.xlsx) using openpyxl and returns base64-encoded strings.

Governed by:
  - RULES_v5.md §A12 (POST /export/excel endpoint)
  - files (4)/PurityForSugar.bas (VBA Base64Decode & file save)
"""

from __future__ import annotations
import io
import json
import base64
from typing import Dict, Any, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_excel_report_base64(data: Union[str, Dict[str, Any]]) -> str:
    """Generates a professional engineering spreadsheet and returns it base64-encoded."""
    if isinstance(data, str):
        try:
            payload = json.loads(data)
        except Exception:
            payload = {}
    else:
        payload = data

    wb = openpyxl.Workbook()

    # Style definitions
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    accent_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    kpi_title_font = Font(name="Calibri", size=10, bold=True, color="595959")
    kpi_val_font = Font(name="Calibri", size=16, bold=True, color="1F497D")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # 1. Summary Sheet
    ws_summary = wb.active
    ws_summary.title = "Balance Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary["A1"] = "PURITY FOR SUGAR — SIMULATION BALANCE REPORT"
    ws_summary["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    ws_summary.merge_cells("A1:G1")

    status = str(payload.get("status", "N/A")).upper()
    iterations = payload.get("iterations", 0)
    final_error = payload.get("final_error", 0.0)

    ws_summary["A3"] = "Simulation Status"
    ws_summary["A3"].font = kpi_title_font
    ws_summary["A4"] = status
    ws_summary["A4"].font = kpi_val_font

    ws_summary["C3"] = "Iterations"
    ws_summary["C3"].font = kpi_title_font
    ws_summary["C4"] = str(iterations)
    ws_summary["C4"].font = kpi_val_font

    ws_summary["E3"] = "Final Relative Error"
    ws_summary["E3"].font = kpi_title_font
    ws_summary["E4"] = f"{final_error:.6e}"
    ws_summary["E4"].font = kpi_val_font

    # Global balance closures if available
    bs = payload.get("balance_summary") or {}
    ws_summary["A6"] = "Global Material Balance Closures"
    ws_summary["A6"].font = Font(name="Calibri", size=12, bold=True, color="1F497D")

    closure_headers = ["Parameter", "Inflow (kg/h)", "Outflow (kg/h)", "Closure Error (%)"]
    for col_idx, h in enumerate(closure_headers, 1):
        cell = ws_summary.cell(row=7, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    rows_data = [
        ["Total Mass Flow", bs.get("total_mass_in_kgh", 0.0), bs.get("total_mass_out_kgh", 0.0), bs.get("mass_closure_error_pct", 0.0)],
        ["Dry Substance (DS)", bs.get("total_ds_in_kgh", 0.0), bs.get("total_ds_out_kgh", 0.0), bs.get("ds_closure_error_pct", 0.0)],
    ]
    for r_idx, r_val in enumerate(rows_data, 8):
        for c_idx, val in enumerate(r_val, 1):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if c_idx > 1:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")

    # 2. Streams Inventory Sheet
    ws_streams = wb.create_sheet(title="Stream Inventory")
    ws_streams.views.sheetView[0].showGridLines = True

    stream_cols = [
        ("Stream ID", "@"),
        ("Mass Flow (kg/h)", "#,##0.0"),
        ("Temp (°C)", "0.0"),
        ("Pressure (kPa)", "0.00"),
        ("DS (%)", "0.00"),
        ("Purity (%)", "0.00"),
        ("Crystals (%)", "0.00"),
        ("TDM (%)", "0.00"),
        ("Enthalpy (kJ/kg)", "#,##0.0"),
        ("Density (kg/m³)", "#,##0.0"),
        ("Fluid Type", "@"),
        ("Water (w/w)", "0.0000"),
        ("Dissolved Sucrose", "0.0000"),
        ("Non-Sucrose #1", "0.0000"),
        ("Non-Sucrose #2", "0.0000"),
        ("Sucrose Crystals", "0.0000"),
        ("Color (ICU)", "#,##0"),
    ]

    for col_idx, (col_name, _) in enumerate(stream_cols, 1):
        c = ws_streams.cell(row=1, column=col_idx, value=col_name)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    flows = payload.get("flows", [])
    for row_idx, f in enumerate(flows, 2):
        row_vals = [
            f.get("id", ""),
            f.get("mass_flow_kgh", 0.0),
            f.get("temperature_c", 0.0),
            f.get("pressure_kpa", 0.0),
            f.get("ds_pct", 0.0),
            f.get("purity_pct", 0.0),
            f.get("crystal_pct", 0.0),
            f.get("tdm_pct", 0.0),
            f.get("enthalpy_kjkg", 0.0),
            f.get("density_kgm3", 0.0) or 0.0,
            f.get("fluid_type", ""),
            f.get("water", 0.0),
            f.get("dissolved_sucrose", 0.0),
            f.get("non_sucrose_1", 0.0),
            f.get("non_sucrose_2", 0.0),
            f.get("sucrose_crystals", 0.0),
            f.get("color_icu", 0.0),
        ]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws_streams.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            fmt = stream_cols[col_idx - 1][1]
            cell.number_format = fmt
            if fmt == "@":
                cell.alignment = Alignment(horizontal="left")
            else:
                cell.alignment = Alignment(horizontal="right")

    # 3. Stations Sheet
    ws_stations = wb.create_sheet(title="Stations Summary")
    ws_stations.views.sheetView[0].showGridLines = True

    station_headers = ["Number", "Name", "Type", "ID", "Calculated Key Indicators"]
    for col_idx, h in enumerate(station_headers, 1):
        c = ws_stations.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center")

    stations = payload.get("stations", [])
    for row_idx, s in enumerate(stations, 2):
        calc_str = ", ".join(f"{k}: {v}" for k, v in s.get("calculated_properties", {}).items())
        r_data = [
            s.get("station_number", 0),
            s.get("name", ""),
            s.get("type", ""),
            s.get("id", ""),
            calc_str,
        ]
        for col_idx, val in enumerate(r_data, 1):
            cell = ws_stations.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")

    # 4. Net Process Revenues Sheet
    rev_data = payload.get("revenues") or payload.get("balance_summary", {}).get("revenues")
    all_sheets = [ws_summary, ws_streams, ws_stations]

    if rev_data and (rev_data.get("inlet_flows") or rev_data.get("outlet_flows") or rev_data.get("total_inlet_cost_per_hour", 0) > 0 or rev_data.get("total_outlet_revenue_per_hour", 0) > 0):
        ws_rev = wb.create_sheet(title="Process Net Revenues")
        ws_rev.views.sheetView[0].showGridLines = True
        all_sheets.append(ws_rev)

        curr = rev_data.get("currency_symbol", "$")
        days = rev_data.get("campaign_days", 300.0)

        ws_rev["A1"] = "PROCESS NET REVENUES REPORT"
        ws_rev["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F497D")
        ws_rev.merge_cells("A1:G1")

        ws_rev["A3"] = "Total Inlet Costs / Day"
        ws_rev["A3"].font = kpi_title_font
        ws_rev["A4"] = f"{curr} {rev_data.get('total_inlet_cost_per_day', 0.0):,.2f}"
        ws_rev["A4"].font = Font(name="Calibri", size=14, bold=True, color="C00000")

        ws_rev["C3"] = "Total Outlet Revenues / Day"
        ws_rev["C3"].font = kpi_title_font
        ws_rev["C4"] = f"{curr} {rev_data.get('total_outlet_revenue_per_day', 0.0):,.2f}"
        ws_rev["C4"].font = Font(name="Calibri", size=14, bold=True, color="008000")

        ws_rev["E3"] = "Net Process Revenues / Day"
        ws_rev["E3"].font = kpi_title_font
        net_day = rev_data.get('net_process_revenue_per_day', 0.0)
        net_col = "008000" if net_day >= 0 else "C00000"
        ws_rev["E4"] = f"{curr} {net_day:,.2f}"
        ws_rev["E4"].font = Font(name="Calibri", size=15, bold=True, color=net_col)

        ws_rev["G3"] = f"Campaign ({days:.0f} days)"
        ws_rev["G3"].font = kpi_title_font
        ws_rev["G4"] = f"{curr} {rev_data.get('net_process_revenue_per_campaign', 0.0):,.2f}"
        ws_rev["G4"].font = Font(name="Calibri", size=14, bold=True, color=net_col)

        curr_row = 6
        # Inlet Costs Table
        ws_rev.cell(row=curr_row, column=1, value="EXTERNAL INLET FLOWS (COSTS)").font = Font(name="Calibri", size=12, bold=True, color="1F497D")
        curr_row += 1

        cost_headers = ["Flow ID", "Flow Name", "Flow (kg/h)", f"Unit Cost ({curr}/kg)", f"Cost / Hour ({curr})", f"Cost / Day ({curr})", f"Campaign Cost ({curr})"]
        for col_idx, h in enumerate(cost_headers, 1):
            c = ws_rev.cell(row=curr_row, column=col_idx, value=h)
            c.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
            c.font = header_font
            c.alignment = Alignment(horizontal="center")
        curr_row += 1

        inlets = rev_data.get("inlet_flows", [])
        if not inlets:
            ws_rev.cell(row=curr_row, column=1, value="No external inlet cost entries assigned.").font = Font(italic=True)
            curr_row += 1
        else:
            for item in inlets:
                r_vals = [
                    item.get("id", ""),
                    item.get("name", ""),
                    item.get("mass_flow_kgh", 0.0),
                    item.get("unit_price", 0.0),
                    item.get("rate_per_hour", 0.0),
                    item.get("rate_per_day", 0.0),
                    item.get("rate_per_campaign", 0.0),
                ]
                for col_idx, val in enumerate(r_vals, 1):
                    cell = ws_rev.cell(row=curr_row, column=col_idx, value=val)
                    cell.border = thin_border
                    if col_idx in (3, 4, 5, 6, 7):
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(horizontal="right")
                curr_row += 1

        curr_row += 2
        # Outlet Revenues Table
        ws_rev.cell(row=curr_row, column=1, value="LEAVING PRODUCT FLOWS (REVENUES)").font = Font(name="Calibri", size=12, bold=True, color="1F497D")
        curr_row += 1

        rev_headers = ["Flow ID", "Flow Name", "Flow (kg/h)", f"Unit Value ({curr}/kg)", f"Revenue / Hour ({curr})", f"Revenue / Day ({curr})", f"Campaign Revenue ({curr})"]
        for col_idx, h in enumerate(rev_headers, 1):
            c = ws_rev.cell(row=curr_row, column=col_idx, value=h)
            c.fill = PatternFill(start_color="008000", end_color="008000", fill_type="solid")
            c.font = header_font
            c.alignment = Alignment(horizontal="center")
        curr_row += 1

        outlets = rev_data.get("outlet_flows", [])
        if not outlets:
            ws_rev.cell(row=curr_row, column=1, value="No product outlet value entries assigned.").font = Font(italic=True)
            curr_row += 1
        else:
            for item in outlets:
                r_vals = [
                    item.get("id", ""),
                    item.get("name", ""),
                    item.get("mass_flow_kgh", 0.0),
                    item.get("unit_price", 0.0),
                    item.get("rate_per_hour", 0.0),
                    item.get("rate_per_day", 0.0),
                    item.get("rate_per_campaign", 0.0),
                ]
                for col_idx, val in enumerate(r_vals, 1):
                    cell = ws_rev.cell(row=curr_row, column=col_idx, value=val)
                    cell.border = thin_border
                    if col_idx in (3, 4, 5, 6, 7):
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(horizontal="right")
                curr_row += 1

    # Auto-adjust column widths
    for sheet in all_sheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)

    # Base64 encode
    return base64.b64encode(buf.read()).decode("ascii")
